import pandas as pd
import re
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.formatting.rule import CellIsRule

FOLDER = Path(__file__).resolve().parent

HOUSE_FILE = FOLDER / "Prince_Georges_Vacancy_Leads_HOUSE_VERIFIED.xlsx"
REDFIN_FILE = FOLDER / "Redfin_All_Downloads_Combined.xlsx"
OUTPUT_FILE = FOLDER / "Prince_Georges_Houses_For_Sale_FINAL_WITH_CONDITION.xlsx"

def clean_col(name):
    return re.sub(r"[^a-z0-9]+", " ", str(name).strip().lower()).strip()

def find_col(df, candidates):
    lookup = {clean_col(c): c for c in df.columns}
    for candidate in candidates:
        key = clean_col(candidate)
        if key in lookup:
            return lookup[key]
    for c in df.columns:
        cc = clean_col(c)
        for candidate in candidates:
            if clean_col(candidate) in cc:
                return c
    return None

def normalize_address(value):
    if pd.isna(value):
        return ""
    s = str(value).upper().strip()
    s = s.replace("&", " AND ")
    s = re.sub(r"[.,#'’]", " ", s)
    s = re.sub(r"\s+", " ", s)

    replacements = {
        r"\bSTREET\b": "ST", r"\bROAD\b": "RD", r"\bAVENUE\b": "AVE",
        r"\bDRIVE\b": "DR", r"\bLANE\b": "LN", r"\bCOURT\b": "CT",
        r"\bCIRCLE\b": "CIR", r"\bBOULEVARD\b": "BLVD",
        r"\bHIGHWAY\b": "HWY", r"\bPARKWAY\b": "PKWY",
        r"\bPLACE\b": "PL", r"\bTERRACE\b": "TER", r"\bTRAIL\b": "TRL",
        r"\bNORTH\b": "N", r"\bSOUTH\b": "S", r"\bEAST\b": "E",
        r"\bWEST\b": "W"
    }
    for pattern, replacement in replacements.items():
        s = re.sub(pattern, replacement, s)
    return s.strip()

print("Loading House-Verified properties...")
houses = pd.read_excel(HOUSE_FILE)

print("Loading combined Redfin listings...")
redfin = pd.read_excel(REDFIN_FILE)

house_address_col = find_col(houses, ["Address", "Property Address", "Street Address"])
redfin_address_col = find_col(redfin, ["Property Address", "Address", "Street Address"])
redfin_status_col = find_col(redfin, ["Status", "Listing Status"])

if not house_address_col:
    raise ValueError("Could not find Address in the House-Verified workbook.")
if not redfin_address_col:
    raise ValueError("Could not find Property Address in the Redfin workbook.")

if redfin_status_col:
    status = redfin[redfin_status_col].fillna("").astype(str).str.strip().str.upper()
    active = redfin[status.eq("ACTIVE")].copy()
else:
    active = redfin.copy()

houses["_match_address"] = houses[house_address_col].map(normalize_address)
active["_match_address"] = active[redfin_address_col].map(normalize_address)

houses = houses[houses["_match_address"] != ""].copy()
active = active[active["_match_address"] != ""].copy()
active = active.drop_duplicates(subset=["_match_address"], keep="first")

matched = houses.merge(
    active,
    on="_match_address",
    how="inner",
    suffixes=("", " - Redfin")
)

matched = matched.drop(columns=["_match_address"], errors="ignore")

# Preserve the automatically calculated score as the starting score.
base_score_col = "Base Vacancy Score"
if "Vacancy Score" in matched.columns:
    matched[base_score_col] = pd.to_numeric(
        matched["Vacancy Score"], errors="coerce"
    ).fillna(0)
else:
    matched[base_score_col] = 0

# Manual inspection columns. These are intentionally blank so the user can inspect each house.
matched["Condition"] = ""
matched["Boarded Up?"] = ""
matched["Overgrown?"] = ""

# Initial score is the automatic/base score. Excel formulas will add manual condition points.
matched["Vacancy Score"] = matched[base_score_col].clip(lower=0, upper=100)

preferred = [
    "Account ID",
    house_address_col,
    "House Verified",
    "Condition",
    "Boarded Up?",
    "Overgrown?",
    base_score_col,
    "Vacancy Score",
    "Priority",
    redfin_status_col,
    find_col(active, ["Price", "List Price", "Listing Price"]),
    find_col(active, ["Beds", "Bedrooms"]),
    find_col(active, ["Baths", "Bathrooms"]),
    find_col(active, ["Square Feet", "Sq Ft", "Square Footage"]),
    find_col(active, ["Year Built"]),
    find_col(active, ["Days On Redfin", "Days On Market"]),
    find_col(active, ["URL", "Listing URL"]),
]

preferred = [c for c in preferred if c and c in matched.columns]
remaining = [c for c in matched.columns if c not in preferred]
matched = matched[preferred + remaining]

matched.to_excel(OUTPUT_FILE, index=False)

# ---------------------------------------------------------
# Excel formatting and live scoring
# ---------------------------------------------------------
wb = load_workbook(OUTPUT_FILE)
ws = wb.active
ws.title = "For Sale Houses"

headers = {cell.value: cell.column for cell in ws[1]}

def col_letter(name):
    return ws.cell(row=1, column=headers[name]).column_letter

condition_col = col_letter("Condition")
boarded_col = col_letter("Boarded Up?")
overgrown_col = col_letter("Overgrown?")
base_col = col_letter(base_score_col)
score_col = col_letter("Vacancy Score")
priority_col = col_letter("Priority")

max_row = ws.max_row

# Dropdown: Condition
condition_validation = DataValidation(
    type="list",
    formula1='"Good,Fair,Poor,Very Poor"',
    allow_blank=True
)
condition_validation.promptTitle = "Property Condition"
condition_validation.prompt = "Choose Good, Fair, Poor, or Very Poor."
condition_validation.errorTitle = "Invalid condition"
condition_validation.error = "Choose a value from the dropdown."
ws.add_data_validation(condition_validation)

# Dropdown: Yes/No
yes_no_validation = DataValidation(
    type="list",
    formula1='"Yes,No"',
    allow_blank=True
)
yes_no_validation.promptTitle = "Yes or No"
yes_no_validation.prompt = "Choose Yes or No."
yes_no_validation.errorTitle = "Invalid choice"
yes_no_validation.error = "Choose Yes or No."
ws.add_data_validation(yes_no_validation)

if max_row >= 2:
    condition_validation.add(f"{condition_col}2:{condition_col}{max_row}")
    yes_no_validation.add(f"{boarded_col}2:{boarded_col}{max_row}")
    yes_no_validation.add(f"{overgrown_col}2:{overgrown_col}{max_row}")

    # Live Vacancy Score:
    # Good = 0, Fair = 10, Poor = 20, Very Poor = 30
    # Boarded Up = +30
    # Overgrown = +30
    for row in range(2, max_row + 1):
        ws.cell(row=row, column=headers["Vacancy Score"]).value = (
            f'=MIN(100,{base_col}{row}+'
            f'IF({condition_col}{row}="Good",0,'
            f'IF({condition_col}{row}="Fair",10,'
            f'IF({condition_col}{row}="Poor",20,'
            f'IF({condition_col}{row}="Very Poor",30,0))))+'
            f'IF({boarded_col}{row}="Yes",30,0)+'
            f'IF({overgrown_col}{row}="Yes",30,0))'
        )

        # Priority updates automatically with the score.
        ws.cell(row=row, column=headers["Priority"]).value = (
            f'=IF({score_col}{row}>=70,"HIGH",'
            f'IF({score_col}{row}>=50,"MEDIUM",'
            f'IF({score_col}{row}>=30,"LOW","VERY LOW")))'
        )

# Add a small instructions sheet.
if "Instructions" in wb.sheetnames:
    del wb["Instructions"]

info = wb.create_sheet("Instructions", 0)
instructions = [
    ["VACANCY SCORE — MANUAL CONDITION CHECKS", ""],
    ["Condition", "Points"],
    ["Good", "0"],
    ["Fair", "10"],
    ["Poor", "20"],
    ["Very Poor", "30"],
    ["Boarded Up?", "Yes = +30 / No = 0"],
    ["Overgrown?", "Yes = +30 / No = 0"],
    ["", ""],
    ["How to use", "Select the dropdown value for each property."],
    ["Vacancy Score", "Updates automatically and is capped at 100."],
    ["Priority", "Updates automatically from the Vacancy Score."],
    ["", ""],
    ["Important", "Condition, boarded windows, and vegetation are manual observations. Do not mark them unless you have inspected the property or have reliable evidence."]
]

for r, row_data in enumerate(instructions, start=1):
    for c, value in enumerate(row_data, start=1):
        info.cell(row=r, column=c).value = value

info["A1"].font = Font(bold=True, size=14)
info["A2"].font = Font(bold=True)
info["B2"].font = Font(bold=True)
info.column_dimensions["A"].width = 85
info.column_dimensions["B"].width = 75

# Format main sheet.
ws.freeze_panes = "A2"
ws.auto_filter.ref = ws.dimensions

for cell in ws[1]:
    cell.font = Font(bold=True)
    cell.alignment = Alignment(horizontal="center", vertical="center")

for name in ["Condition", "Boarded Up?", "Overgrown?"]:
    ws.column_dimensions[col_letter(name)].width = 18

ws.column_dimensions[base_col].width = 18
ws.column_dimensions[score_col].width = 15
ws.column_dimensions[priority_col].width = 12

# Highlight score cells based on value.
score_range = f"{score_col}2:{score_col}{max_row}"
ws.conditional_formatting.add(
    score_range,
    CellIsRule(operator="greaterThanOrEqual", formula=["70"], fill=PatternFill("solid", fgColor="F4CCCC"))
)
ws.conditional_formatting.add(
    score_range,
    CellIsRule(operator="between", formula=["50", "69"], fill=PatternFill("solid", fgColor="FFF2CC"))
)
ws.conditional_formatting.add(
    score_range,
    CellIsRule(operator="lessThan", formula=["50"], fill=PatternFill("solid", fgColor="D9EAD3"))
)

wb.calculation.fullCalcOnLoad = True
wb.calculation.forceFullCalc = True
wb.save(OUTPUT_FILE)

print()
print("===================================================")
print("FINAL PROGRAM CREATED SUCCESSFULLY")
print("===================================================")
print(f"Matching properties: {len(matched):,}")
print(f"Created: {OUTPUT_FILE.name}")
print()
print("Manual columns added:")
print("  Condition: Good / Fair / Poor / Very Poor")
print("  Boarded Up?: Yes / No")
print("  Overgrown?: Yes / No")
print()
print("Condition points:")
print("  Good = 0")
print("  Fair = 10")
print("  Poor = 20")
print("  Very Poor = 30")
print("  Boarded Up = +30")
print("  Overgrown = +30")
print()
print("Vacancy Score updates automatically and is capped at 100.")
print("===================================================")
